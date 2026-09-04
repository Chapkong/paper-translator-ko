"""Generate a synthetic English 'paper' (DOCX + PDF) with figures and a table for pipeline testing."""
import io, os
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches

HERE = os.path.dirname(os.path.abspath(__file__))
INPUT = os.path.join(HERE, "..", "input")
os.makedirs(INPUT, exist_ok=True)

def fig(path, kind):
    plt.figure(figsize=(5, 3))
    if kind == 1:
        plt.plot([2018, 2019, 2020, 2021, 2022], [12, 15, 21, 26, 34], marker="o")
        plt.title("Figure 1. Growth of doctoral graduates in STEM")
        plt.xlabel("Year"); plt.ylabel("Thousands")
    else:
        plt.bar(["Academia", "Industry", "Government", "Abroad"], [41, 38, 12, 9])
        plt.title("Figure 2. First-job destination of PhD holders (%)")
    plt.tight_layout(); plt.savefig(path, dpi=120); plt.close()

f1 = os.path.join(INPUT, "_fig1.png"); f2 = os.path.join(INPUT, "_fig2.png")
fig(f1, 1); fig(f2, 2)

doc = Document()
doc.add_heading("Effective Supply of Science and Technology Talent: A Pipeline Perspective", 0)
doc.add_paragraph("Jane Doe, John Smith — Institute for Policy Studies")
doc.add_heading("Abstract", 1)
doc.add_paragraph(
    "Conventional workforce projections count degree holders as if every graduate enters the "
    "domestic research labor market. We argue that this overstates supply. Using a porous-pipeline "
    "framework, we estimate the effective supply of doctoral-level researchers by subtracting "
    "leakage at four transition points: field switching, emigration, career exit, and underemployment. "
    "Applying the framework to administrative data from 2018 to 2022, we find that effective supply "
    "is 27% lower than headcount-based estimates, with the gap widening over time.")
doc.add_heading("1. Introduction", 1)
doc.add_paragraph(
    "Governments routinely justify expansion of graduate programs by pointing to projected shortages "
    "of researchers. However, the link between the number of degrees awarded and the number of "
    "researchers actually available to employers is weaker than these projections assume. "
    "This paper addresses that gap by distinguishing nominal supply from effective supply.")
doc.add_paragraph(
    "The remainder of the paper is organized as follows. Section 2 reviews prior work on pipeline "
    "leakage. Section 3 describes the data and estimation strategy. Section 4 presents results, "
    "and Section 5 discusses policy implications.")
doc.add_heading("2. Related Work", 1)
doc.add_paragraph(
    "The 'leaky pipeline' metaphor originated in studies of gender attrition in academic careers "
    "(Berryman, 1983) and was later generalized to describe any systematic loss of talent between "
    "training and employment. Critics note that the metaphor implies a single linear path, whereas "
    "real careers involve re-entry and lateral moves (Cannady et al., 2014). We adopt the term "
    "'porous pipeline' to acknowledge bidirectional flows.")
doc.add_heading("3. Data and Methods", 1)
doc.add_paragraph(
    "We link three administrative sources: the national graduate survey, employment insurance "
    "records, and the researcher registry. The linked panel covers 48,213 doctoral recipients. "
    "Figure 1 shows the growth in annual doctoral graduates over the study period.")
doc.add_picture(f1, width=Inches(4.5))
doc.add_paragraph("Figure 1. Growth of doctoral graduates in STEM fields, 2018–2022.")
doc.add_paragraph(
    "Effective supply S_eff is defined as S_eff = N × (1 − l_1)(1 − l_2)(1 − l_3)(1 − l_4), "
    "where N is the number of new degree holders and l_k is the leakage rate at transition k. "
    "Leakage rates are estimated with a discrete-time hazard model controlling for field, "
    "institution tier, and cohort.")
doc.add_heading("4. Results", 1)
doc.add_paragraph(
    "Table 1 reports estimated leakage rates. Emigration is the largest single source of loss, "
    "followed by career exit. Figure 2 shows first-job destinations of PhD holders.")
t = doc.add_table(rows=5, cols=3); t.style = "Table Grid"
rows = [("Transition", "Leakage rate", "95% CI"),
        ("Field switching", "6.1%", "5.4–6.8"),
        ("Emigration", "11.3%", "10.5–12.1"),
        ("Career exit", "8.7%", "8.0–9.4"),
        ("Underemployment", "4.2%", "3.7–4.7")]
for i, r in enumerate(rows):
    for j, c in enumerate(r):
        t.cell(i, j).text = c
doc.add_paragraph("Table 1. Estimated leakage rates by transition point.")
doc.add_picture(f2, width=Inches(4.5))
doc.add_paragraph("Figure 2. First-job destination of PhD holders (%).")
doc.add_heading("5. Discussion and Policy Implications", 1)
doc.add_paragraph(
    "Our findings suggest that policies aimed solely at increasing enrollment will have a muted "
    "effect on the research workforce unless leakage is addressed simultaneously. Retention "
    "measures—competitive early-career salaries, stable positions, and return incentives for "
    "researchers abroad—may yield higher marginal returns than further expansion of doctoral programs.")
doc.add_heading("References", 1)
doc.add_paragraph("Berryman, S. (1983). Who Will Do Science? Rockefeller Foundation.")
doc.add_paragraph("Cannady, M., Greenwald, E., & Harris, K. (2014). Problematizing the STEM pipeline metaphor. Science Education, 98(3), 443–460.")

docx_path = os.path.join(INPUT, "sample.docx")
doc.save(docx_path)
os.system(f'cd "{INPUT}" && soffice --headless --convert-to pdf sample.docx >/dev/null 2>&1')
for f in (f1, f2):
    os.remove(f)
print("wrote", docx_path, "and sample.pdf")
